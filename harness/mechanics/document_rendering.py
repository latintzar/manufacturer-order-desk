"""Bounded document creation from text, using standard document libraries.

The worker accepts source text, not uploaded Office files, templates or Python.
Markdown links and image references remain text; rendering performs no fetches.
"""
from __future__ import annotations

import csv
from io import BytesIO, StringIO
import json
import re
from pathlib import Path
import subprocess
import sys
import threading

MAX_SOURCE_BYTES = 128 * 1024
MAX_OUTPUT_BYTES = 8 * 1024 * 1024
MAX_BLOCKS = 1000
MAX_ROWS = 2000
MAX_COLUMNS = 100
MAX_CELLS = 20000
_CONVERTERS = threading.BoundedSemaphore(1)


def render(suffix, title, content):
    if len(content.encode('utf-8')) > MAX_SOURCE_BYTES:
        raise ValueError('Split formatted documents into source text of up to 128 KiB each.')
    if any(ord(char) < 32 and char not in '\t\n\r' for char in content):
        raise ValueError('Use readable text and standard line breaks for this document.')
    if not _CONVERTERS.acquire(timeout=1):
        raise ValueError('A document is being prepared. Try this file again shortly.')
    try:
        try:
            result = subprocess.run([sys.executable, str(Path(__file__).resolve()), '--render'],
                input=json.dumps({'suffix': suffix, 'title': title, 'content': content}).encode(),
                stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=15, check=False,
                env={'PYTHONDONTWRITEBYTECODE': '1', 'OPENBLAS_NUM_THREADS': '1', 'OMP_NUM_THREADS': '1'}, close_fds=True)
        except subprocess.TimeoutExpired:
            raise ValueError('Split this document into smaller sections and try again.') from None
        if result.returncode:
            # The worker emits only controlled errors; tracebacks and content stay private.
            if result.returncode == 2:
                raise ValueError(result.stderr.decode('utf-8', errors='replace')[:300])
            raise ValueError('This document needs a smaller section or a simpler table.')
        if not result.stdout or len(result.stdout) > MAX_OUTPUT_BYTES:
            raise ValueError('Split this document into smaller files.')
        return result.stdout
    finally:
        _CONVERTERS.release()


def _inline(tokens):
    runs, bold, italic, links = [], False, False, []
    for token in tokens or []:
        kind = token.type
        if kind in {'strong_open', 'strong_close'}: bold = kind.endswith('open')
        elif kind in {'em_open', 'em_close'}: italic = kind.endswith('open')
        elif kind == 'link_open': links.append(token.attrGet('href') or '')
        elif kind == 'link_close':
            target = links.pop() if links else ''
            if target: runs.append({'text': ' (' + target + ')', 'bold': False, 'italic': False})
        elif kind in {'text', 'code_inline', 'softbreak', 'hardbreak', 'image', 'html_inline'}:
            text = '\n' if kind in {'softbreak', 'hardbreak'} else token.content
            if kind == 'image': text = '[Image: ' + text + ']'
            runs.append({'text': text, 'bold': bold, 'italic': italic, 'code': kind == 'code_inline'})
    return runs


def _blocks(content):
    from markdown_it import MarkdownIt
    parser = MarkdownIt('commonmark', {'html': False, 'maxNesting': 20}).enable('table')
    tokens = parser.parse(content)
    if len(tokens) > MAX_BLOCKS * 8:
        raise ValueError('Split this document into fewer than 1,000 paragraphs and table rows.')
    blocks, lists, quote, index = [], [], 0, 0
    while index < len(tokens):
        token = tokens[index]
        if token.type == 'table_open':
            rows, row = [], []
            index += 1
            while index < len(tokens) and tokens[index].type != 'table_close':
                current = tokens[index]
                if current.type == 'tr_open': row = []
                elif current.type == 'inline': row.append(_inline(current.children))
                elif current.type == 'tr_close': rows.append(row)
                index += 1
            if len(rows) > 100 or any(len(row) > 12 for row in rows):
                raise ValueError('Use up to 100 rows and 12 columns in a document table, or create an XLSX file.')
            blocks.append({'kind': 'table', 'rows': rows})
        elif token.type in {'bullet_list_open', 'ordered_list_open'}:
            lists.append({'ordered': token.type == 'ordered_list_open', 'number': int(token.attrGet('start') or 1) - 1})
        elif token.type in {'bullet_list_close', 'ordered_list_close'}:
            if lists: lists.pop()
        elif token.type == 'list_item_open' and lists:
            lists[-1]['number'] += 1
        elif token.type == 'blockquote_open': quote += 1
        elif token.type == 'blockquote_close': quote -= 1
        elif token.type == 'inline':
            previous = tokens[index - 1]
            heading = int(previous.tag[1:]) if previous.type == 'heading_open' else 0
            prefix = (str(lists[-1]['number']) + '. ' if lists[-1]['ordered'] else '• ') if lists else ''
            blocks.append({'kind': 'heading' if heading else 'paragraph', 'level': heading,
                           'runs': _inline(token.children), 'prefix': prefix, 'indent': min(len(lists), 4), 'quote': bool(quote)})
        elif token.type in {'code_block', 'fence'}:
            blocks.append({'kind': 'code', 'runs': [{'text': token.content.rstrip('\n')}]})
        elif token.type == 'hr':
            blocks.append({'kind': 'paragraph', 'runs': [{'text': '—'}]})
        index += 1
    if len(blocks) > MAX_BLOCKS:
        raise ValueError('Split this document into fewer than 1,000 paragraphs.')
    return blocks


def _docx(title, content):
    from docx import Document
    from docx.shared import Mm, Pt
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(20)
    section.left_margin = section.right_margin = Mm(22)
    document.core_properties.title, document.core_properties.author = title, 'Rocoto'
    normal = document.styles['Normal']
    normal.font.name, normal.font.size = 'Arial', Pt(11)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(8)
    for level in range(1, 7):
        style = document.styles['Heading ' + str(level)]
        style.font.name, style.font.size = 'Arial', Pt({1: 22, 2: 17, 3: 13}.get(level, 11))
        style.paragraph_format.keep_with_next = True
    def add_runs(paragraph, runs):
        for item in runs:
            run = paragraph.add_run(item['text'])
            run.bold, run.italic = item.get('bold', False), item.get('italic', False)
            if item.get('code'): run.font.name = 'Courier New'
    for block in _blocks(content):
        if block['kind'] == 'table':
            rows = block['rows']
            if not rows: continue
            table = document.add_table(rows=0, cols=max(map(len, rows)))
            table.style = 'Light Shading Accent 1'
            for row in rows:
                for cell, runs in zip(table.add_row().cells, row): add_runs(cell.paragraphs[0], runs)
        else:
            paragraph = document.add_paragraph(style='Heading ' + str(block['level']) if block['kind'] == 'heading' else None)
            paragraph.paragraph_format.left_indent = Mm(block.get('indent', 0) * 5)
            if block.get('quote'): paragraph.paragraph_format.left_indent = Mm(8)
            if block.get('prefix'): paragraph.add_run(block['prefix'])
            add_runs(paragraph, block['runs'])
            if block['kind'] == 'code':
                for run in paragraph.runs: run.font.name, run.font.size = 'Courier New', Pt(9)
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _pdf(title, content):
    from html import escape
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle
    from reportlab.platypus.doctemplate import LayoutError
    fonts = Path('/usr/share/fonts/truetype/dejavu')
    for name, filename in [('RocotoSans', 'DejaVuSans.ttf'), ('RocotoSansBold', 'DejaVuSans-Bold.ttf')]:
        if not (fonts / filename).is_file(): raise ValueError('PDF creation needs the configured DejaVu fonts.')
        pdfmetrics.registerFont(TTFont(name, str(fonts / filename)))
    pdfmetrics.registerFontFamily('RocotoSans', normal='RocotoSans', bold='RocotoSansBold', italic='RocotoSans', boldItalic='RocotoSansBold')
    glyphs = pdfmetrics.getFont('RocotoSans').face.charToGlyph
    if any(ord(char) not in glyphs and char not in '\r\n\t' for char in content):
        raise ValueError('This PDF font supports Latin, Cyrillic and Greek text. Use DOCX for additional characters.')
    body = ParagraphStyle('Body', fontName='RocotoSans', fontSize=10.5, leading=15, spaceAfter=8)
    heading = {level: ParagraphStyle('Heading' + str(level), parent=body, fontName='RocotoSansBold',
        fontSize={1: 22, 2: 16, 3: 13}.get(level, 11), leading={1: 28, 2: 21, 3: 18}.get(level, 16),
        spaceBefore=12, spaceAfter=8, keepWithNext=True) for level in range(1, 7)}
    def text(runs):
        pieces = []
        for run in runs:
            part = escape(run['text']).replace('\n', '<br/>')
            if run.get('bold'): part = '<b>' + part + '</b>'
            if run.get('italic'): part = '<i>' + part + '</i>'
            pieces.append(part)
        return ''.join(pieces)
    stream, story = BytesIO(), []
    for block in _blocks(content):
        if block['kind'] == 'table':
            rows = [[Paragraph(text(cell), body) for cell in row] for row in block['rows']]
            if not rows: continue
            width = max(map(len, rows))
            table = Table(rows, colWidths=[(A4[0] - 44 * mm) / width] * width, repeatRows=1)
            table.setStyle(TableStyle([('VALIGN', (0, 0), (-1, -1), 'TOP'),
                ('LINEBELOW', (0, 0), (-1, 0), 0.5, colors.grey), ('BOTTOMPADDING', (0, 0), (-1, -1), 7)]))
            story.append(table)
        else:
            style = heading[block['level']] if block['kind'] == 'heading' else ParagraphStyle('Paragraph', parent=body,
                leftIndent=block.get('indent', 0) * 5 * mm + (8 * mm if block.get('quote') else 0))
            story.append(Paragraph(escape(block.get('prefix', '')) + text(block['runs']) or ' ', style))
    def footer(canvas, document):
        if document.page > 100: raise ValueError('Split this PDF into documents of up to 100 pages.')
        canvas.saveState()
        canvas.setFont('RocotoSans', 9)
        canvas.setFillColor(colors.grey)
        canvas.drawRightString(A4[0] - 22 * mm, 12 * mm, str(document.page))
        canvas.restoreState()
    try:
        SimpleDocTemplate(stream, pagesize=A4, leftMargin=22*mm, rightMargin=22*mm,
            topMargin=20*mm, bottomMargin=20*mm, title=title, author='Rocoto').build(story or [Paragraph(' ', body)],
            onFirstPage=footer, onLaterPages=footer)
    except LayoutError:
        raise ValueError('Split large table cells into shorter paragraphs or create an XLSX file.') from None
    return stream.getvalue()


def _xlsx(title, content):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    try:
        delimiter = csv.Sniffer().sniff(content[:8192], delimiters=',\t').delimiter
    except csv.Error:
        delimiter = '\t' if '\t' in content.partition('\n')[0] else ','
    rows = list(csv.reader(StringIO(content), delimiter=delimiter, strict=True))
    if (len(rows) > MAX_ROWS or sum(map(len, rows)) > MAX_CELLS
            or any(len(row) > MAX_COLUMNS for row in rows)):
        raise ValueError('Use up to 2,000 rows, 100 columns and 20,000 cells per spreadsheet.')
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Data'
    workbook.properties.title, workbook.properties.creator = title, 'Rocoto'
    lengths = {}
    for row_index, row in enumerate(rows, 1):
        for column_index, value in enumerate(row, 1):
            if len(value) > 32767:
                raise ValueError('Use spreadsheet cell text of up to 32,767 characters.')
            cell = sheet.cell(row_index, column_index)
            # All input is data. An equals sign must remain text rather than become
            # a formula, link, macro or external workbook reference when opened.
            cell.value, cell.data_type = value, 's'
            if (re.fullmatch(r'-?(?:0|[1-9]\d*)(?:\.\d+)?', value)
                    and len(value.replace('-', '').replace('.', '')) <= 15):
                cell.value, cell.data_type = (float(value) if '.' in value else int(value)), 'n'
            cell.font = Font(name='Arial', size=11, bold=row_index == 1)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if row_index == 1: cell.fill = PatternFill('solid', fgColor='EEF0F3')
            lengths[column_index] = max(lengths.get(column_index, 0), min(55, max(map(len, value.splitlines() or ['']))))
    for column_index, length in lengths.items():
        sheet.column_dimensions[get_column_letter(column_index)].width = max(14, length + 3)
    sheet.freeze_panes = 'A2'
    if rows: sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _worker():
    import resource
    resource.setrlimit(resource.RLIMIT_AS, (256 * 1024 * 1024, 256 * 1024 * 1024))
    resource.setrlimit(resource.RLIMIT_CPU, (8, 10))
    resource.setrlimit(resource.RLIMIT_NOFILE, (32, 32))
    resource.setrlimit(resource.RLIMIT_FSIZE, (MAX_OUTPUT_BYTES, MAX_OUTPUT_BYTES))
    document = json.loads(sys.stdin.buffer.read(MAX_SOURCE_BYTES * 7 + 4096))
    result = {'.docx': _docx, '.pdf': _pdf, '.xlsx': _xlsx}[document['suffix']](document['title'], document['content'])
    if len(result) > MAX_OUTPUT_BYTES: raise ValueError('Split this document into smaller files.')
    sys.stdout.buffer.write(result)


if __name__ == '__main__':
    try:
        _worker()
    except (ValueError, csv.Error) as exc:
        sys.stderr.write(str(exc)[:300])
        sys.exit(2)
    except Exception:
        sys.exit(1)
