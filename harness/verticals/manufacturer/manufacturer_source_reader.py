"""Bounded subprocess for inert source-document extraction; never imports the app."""
from __future__ import annotations

import csv
import io
import json
import resource
import sys
import zipfile

MAX_TEXT = 240000
MAX_ZIP = 64 * 1024 * 1024


def read(raw, suffix):
    parts, pages, partial = [], 0, False
    if suffix in {'.xlsx', '.docx'}:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(item.file_size for item in entries) > MAX_ZIP:
                raise ValueError('archive allowance')
    if suffix == '.pdf':
        import fitz
        with fitz.open(stream=raw, filetype='pdf') as document:
            if document.needs_pass:
                raise ValueError('encrypted')
            pages = len(document)
            partial = pages > 50
            for index in range(min(pages, 50)):
                text = document[index].get_text()
                if not text.strip():
                    partial = True
                    text = '[This page needs a visual source review.]'
                parts.append(f'Page {index + 1}\n{text}')
                if sum(map(len, parts)) > MAX_TEXT:
                    partial = True
                    break
    elif suffix == '.xlsx':
        import openpyxl
        book = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True, keep_links=False)
        try:
            pages = len(book.worksheets)
            partial = pages > 20
            for sheet in book.worksheets[:20]:
                parts.append('Sheet: ' + sheet.title)
                for index, row in enumerate(sheet.iter_rows(values_only=True)):
                    if index >= 10000 or len(row) > 100:
                        partial = True
                        break
                    parts.append('\t'.join('' if cell is None else str(cell) for cell in row))
                    if sum(map(len, parts)) > MAX_TEXT:
                        partial = True
                        break
                if sum(map(len, parts)) > MAX_TEXT:
                    break
        finally:
            book.close()
    elif suffix == '.docx':
        from docx import Document
        document = Document(io.BytesIO(raw))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            parts += ['\t'.join(cell.text for cell in row.cells) for row in table.rows]
    else:
        text = raw.decode('utf-8-sig')
        if '\x00' in text:
            raise ValueError('binary')
        if suffix == '.json':
            json.loads(text)
        elif suffix in {'.csv', '.tsv'}:
            csv.field_size_limit(100000)
            for row in csv.reader(io.StringIO(text), delimiter='\t' if suffix == '.tsv' else ','):
                if len(row) > 1000:
                    raise ValueError('columns')
        parts = [text]
    text = '\n'.join(parts)
    partial = partial or len(text) > MAX_TEXT
    text = text[:MAX_TEXT]
    return {'text': text, 'pages': pages or None, 'characters': len(text),
            'status': 'needs_review' if partial or not text.strip() else 'ready'}


if __name__ == '__main__':
    resource.setrlimit(resource.RLIMIT_AS, (384 * 1024 * 1024, 384 * 1024 * 1024))
    resource.setrlimit(resource.RLIMIT_CPU, (15, 15))
    resource.setrlimit(resource.RLIMIT_FSIZE, (4 * 1024 * 1024, 4 * 1024 * 1024))
    raw = sys.stdin.buffer.read(16 * 1024 * 1024 + 1)
    if len(raw) > 16 * 1024 * 1024:
        raise ValueError('upload allowance')
    if sys.argv[1] == '.pdf-page':
        import fitz
        page = int(sys.argv[2])
        with fitz.open(stream=raw, filetype='pdf') as document:
            if document.needs_pass or not 1 <= page <= min(len(document), 10000):
                raise ValueError('page')
            selected = document[page - 1]
            side = max(selected.rect.width, selected.rect.height)
            if side <= 0:
                raise ValueError('page size')
            image = selected.get_pixmap(matrix=fitz.Matrix(1600 / side, 1600 / side), alpha=False)
            sys.stdout.buffer.write(str(len(document)).encode() + b'\n' + image.tobytes('png'))
    else:
        print(json.dumps(read(raw, sys.argv[1]), ensure_ascii=False))
